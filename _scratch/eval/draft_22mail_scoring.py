#!/usr/bin/env python3
"""Draft ①②③④ + fail_tag for Joyce 22-mail eval · NOT official scores (human must confirm)."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

EVAL = Path(__file__).resolve().parent
ROOT = EVAL.parents[1]
MAP_PATH = EVAL / "cs_email_query_map.json"

# Reference attachment hints in corpus md (## 客服配图 / attachments/)
ATTACHMENT_HINT = re.compile(r"attachments/[^\s`|]+\.(?:png|jpg|jpeg|gif)", re.I)
NUMBERED_STEP = re.compile(r"^\s*\d+\.\s", re.M)

# Reference / reply grounding hints for ② draft (not semantic QA).
_REF_KEY_PATTERNS: list[tuple[str, str]] = [
    (r"11#\s*(?:and\s*)?12#", "11#/12#"),
    (r"4#\s*(?:and\s*)?5#", "4#/5#"),
    (r"dip\s*switch\s*#?\s*3", "DIP#3"),
    (r"dip\s*switch\s*#?\s*5", "DIP#5"),
    (r"instant(?:aneous)?\s+short", "instant_short"),
    (r"force\s+potentiometer", "FORCE"),
    (r"soft\s+stop", "SOFT_STOP"),
    (r"limit\s+switch\s*b", "limit_B"),
    (r"ult,\s*com.*dlt|ulmt.*com.*dlmt", "limit_short"),
    (r"extension\s+cable|short\s+cable", "extension_cable"),
    (r"video|dropbox|google\s+drive", "media_request"),
    (r"shipping\s+address|zip\s+code", "address_request"),
]


def ref_key_hits(text: str) -> set[str]:
    hits: set[str] = set()
    for pat, label in _REF_KEY_PATTERNS:
        if re.search(pat, text, re.IGNORECASE):
            hits.add(label)
    return hits


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def extract_reference_reply(md_text: str) -> str:
    if "## 客服回复" not in md_text:
        return ""
    tail = md_text.split("## 客服回复", 1)[1]
    parts = tail.split("```")
    if len(parts) >= 3:
        return parts[1].strip()
    return ""


def reference_wants_images(md_text: str) -> bool:
    if ATTACHMENT_HINT.search(md_text):
        return True
    if "## 客服配图" in md_text or "配图" in md_text.split("## 手册对照")[0][-200:]:
        block = md_text
        if "## 客服配图" in md_text:
            block = md_text.split("## 客服配图", 1)[1][:800]
        return "attachments/" in block or ".png" in block
    return False


def unique_image_files(images: list) -> list[str]:
    seen: list[str] = []
    for img in images or []:
        name = img.get("file") or "?"
        if name not in seen:
            seen.append(name)
    return seen


def draft_dim1(row: dict, scen: dict) -> tuple[str, str]:
    """① 机型/问题 · reason snippet."""
    mail_type = scen.get("type") or "fault"
    hit = row.get("top1_hit")
    if hit is True:
        return "对", "top1_hit=true"
    if hit is False:
        return "错", f"top1={row.get('top1_group')} not in acceptable"
    if mail_type == "presales":
        style = (row.get("style") or {}).get("family_id") or ""
        if style == "F7_presales" or "topens.com" in (row.get("generated_reply_en") or "").lower():
            return "对", "presales · style/route OK"
        return "待人工", "presales · top1_hit=null"
    exp = scen.get("expected_library")
    matched = row.get("matched_library")
    if exp and matched == exp:
        return "对", f"library={matched}"
    if exp and matched != exp:
        return "错", f"expected {exp} got {matched}"
    return "待人工", "no acceptable_groups · fan_out"


def draft_dim2(row: dict, scen: dict, ref: str) -> tuple[str, str]:
    """② 英文可用度 draft."""
    reply = (row.get("generated_reply_en") or "").strip()
    if not reply:
        return "需重写", "empty reply"
    mail_type = scen.get("type") or "fault"
    audit = row.get("reply_audit") or {}
    notes: list[str] = []

    if row.get("reply_truncated") or audit.get("truncated"):
        return "需重写", "truncated"
    if audit.get("glued_steps"):
        notes.append("glued_steps")
    if audit.get("proceed_ref"):
        notes.append("proceed_ref")

    low = reply.lower()
    if mail_type == "presales":
        thin = (
            "not covered" in low
            or "unfortunately" in low
            or "do not have" in low
            or "not available" in low
        )
        has_links = "topens.com" in low or "amazon" in low
        if thin and not has_links:
            return "需重写", "presales · 拒答/无链接"
        if has_links and len(reply) > 200:
            level = "直接可发" if not notes and not thin else "小改可发"
            return level, "presales + links" + (" · thin" if thin else "")
        return "小改可发", "presales · weak links"

    # DIP #3 vs #5 (AT/A3S qa_011 path)
    top1 = row.get("top1_group") or ""
    lib = row.get("matched_library") or ""
    if top1 == "qa_011" or (lib == "a3s" and "qa_011" in str(row.get("top3_groups"))):
        if re.search(r"dip\s*switch\s*#?\s*5", low) and "dip switch #3" not in low:
            return "小改可发", "qa_011 path but DIP #5 (expect #3 for A3S)"
    if top1 == "qa_010" and re.search(r"dip\s*switch\s*#?\s*3\b", low) and "#5" not in low:
        return "小改可发", "qa_010 path but DIP #3 only"

    steps = audit.get("step_count") or len(NUMBERED_STEP.findall(reply))
    ref_steps = len(NUMBERED_STEP.findall(ref)) if ref else 0
    if ref_steps and steps < max(2, ref_steps - 1):
        notes.append(f"steps={steps} vs ref≈{ref_steps}")

    if notes:
        return "小改可发", "; ".join(notes)

    ref_keys = ref_key_hits(ref) if ref else set()
    reply_keys = ref_key_hits(reply)
    has_gold_ref = bool(ref.strip())

    if ref_keys:
        overlap = len(ref_keys & reply_keys) / len(ref_keys)
        missing = sorted(ref_keys - reply_keys)
        note = f"ref_keys={int(overlap * 100)}%"
        if missing:
            note += f" ·缺 {','.join(missing[:3])}"
        if overlap >= 0.75 and steps >= max(2, ref_steps - 1):
            return "直接可发", note
        if overlap >= 0.5:
            return "小改可发", note + " · 待逐步对照 Reference"
        return "小改可发", note + " · 关键项偏差大"

    if steps >= 3 and "11#" in reply and ("4#" in reply or "4# and 5#" in reply):
        return "直接可发", f"steps={steps} · key terminals present"

    # No ref_keys patterns: do not auto-④否决 — human must confirm (§7.2)
    if row.get("top1_hit") is True and steps >= 3:
        if not has_gold_ref:
            return "待人工", f"top1_hit · steps={steps} · 无金标准回信"
        return "待人工", f"top1_hit · steps={steps} · Reference 无 ref_keys 模式"

    if steps >= 2:
        if not has_gold_ref:
            return "待人工", f"steps={steps} · 无金标准回信"
        return "待人工", f"steps={steps} · Reference 无 ref_keys 模式"

    return "待人工", "short or unnumbered · read Reference"


def draft_dim3(row: dict, md_text: str) -> tuple[str, str]:
    """③ 图片配对 draft."""
    imgs = unique_image_files(row.get("images_used"))
    wants = reference_wants_images(md_text)
    if not wants and not imgs:
        return "无需图", "Reference 未要求配图"
    if wants and not imgs:
        return "该配没配", "Reference 有配图 · images_used 空"
    if not wants and imgs:
        return "配对相关", "有图 · Reference 未强制"
    if len(imgs) > 1 and len(row.get("images_used") or []) > len(imgs):
        return "配对相关", f"dedup={imgs} · batch 重复条目"
    null_cap = sum(
        1 for i in (row.get("images_used") or []) if not (i.get("caption") or i.get("caption_en"))
    )
    if null_cap and imgs:
        return "配对相关", f"{imgs[0]} · caption 空(batch 路径)"
    return "配对相关", ", ".join(imgs)


def draft_dim4(d1: str, d2: str, d3: str) -> str:
    if "待人工" in (d1, d2, d3):
        return "待人工"
    if d1 != "对":
        return "不通过"
    if d2 != "直接可发":
        return "不通过"
    if d3 in ("图不对", "该配没配"):
        return "不通过"
    return "通过"


def fail_tag(d1: str, d2: str, d3: str, row: dict) -> str:
    tags: list[str] = []
    if d1 == "错" or "待人工" in d1:
        tags.append("机型" if d1 == "错" else "机型?")
    if d2 in ("需重写", "小改可发") or "待人工" in d2:
        tags.append("生成")
    if d3 in ("图不对", "该配没配"):
        tags.append("图片")
    if row.get("top1_hit") is False:
        tags.append("检索")
    if len(tags) > 1:
        return "复合"
    return tags[0] if tags else "—"


def draft_case(row: dict, scen: dict) -> dict:
    md_path = ROOT / scen["file"] if scen.get("file") else None
    md_text = md_path.read_text(encoding="utf-8") if md_path and md_path.is_file() else ""
    ref = extract_reference_reply(md_text)
    d1, r1 = draft_dim1(row, scen)
    d2, r2 = draft_dim2(row, scen, ref)
    d3, r3 = draft_dim3(row, md_text)
    d4 = draft_dim4(d1, d2, d3)
    return {
        "scenario_id": row["scenario_id"],
        "mail_id": row.get("mail_id"),
        "mvp19": row.get("mvp19"),
        "dim1": d1,
        "dim1_note": r1,
        "dim2": d2,
        "dim2_note": r2,
        "dim3": d3,
        "dim3_note": r3,
        "dim4": d4,
        "fail_tag": fail_tag(d1, d2, d3, row),
        "top1": row.get("top1_group"),
        "style": (row.get("style") or {}).get("family_id"),
    }


def render_md(cases: list[dict], meta: dict, *, title: str) -> str:
    lines = [
        f"# {title}",
        "",
        "**性质**：四维 **草稿** · 须人工确认后写入 xlsx E–H · **不可**作发链依据",
        f"**来源**：`{meta.get('source_json', '')}` · `{meta.get('generated_at', '')}`",
        "",
        "## MVP 19 草稿汇总",
        "",
    ]
    mvp = [c for c in cases if c.get("mvp19") == "是"]
    pass_strict = sum(1 for c in mvp if c["dim4"] == "通过")
    pending = sum(1 for c in mvp if c["dim4"] == "待人工")
    lines.append(f"- MVP 子集 **{len(mvp)}** 封 · 草稿④通过 **{pass_strict}** · 待人工 **{pending}** · 目标 ≥16")
    d2_direct = sum(1 for c in mvp if c["dim2"] == "直接可发")
    d2_minor = sum(1 for c in mvp if c["dim2"] == "小改可发")
    d2_rewrite = sum(1 for c in mvp if c["dim2"] == "需重写")
    lines.append(
        f"- ②草稿（MVP19）：直接可发 **{d2_direct}** · 小改可发 **{d2_minor}** · 需重写 **{d2_rewrite}**"
    )
    lines.append("")
    lines.append("## ② 英文可用度 · MVP19 草稿")
    lines.append("")
    lines.append("| mail | cs_id | ②草稿 | 依据 | top1 |")
    lines.append("| ---: | --- | --- | --- | --- |")
    for c in mvp:
        note = c["dim2_note"]
        lines.append(
            f"| {c.get('mail_id','').replace('MAIL-','#')} | {c['scenario_id']} | "
            f"**{c['dim2']}** | {note} | {c.get('top1','')} |"
        )
    lines.append("")
    lines.append("## 全表（①–④）")
    lines.append("")
    lines.append("| cs_id | MVP19 | ① | ② | ③ | ④ | fail_tag | 备注 |")
    lines.append("| --- | :---: | --- | --- | --- | --- | --- | --- |")
    for c in cases:
        note = c["dim2_note"][:40] + ("…" if len(c["dim2_note"]) > 40 else "")
        lines.append(
            f"| {c['scenario_id']} | {c.get('mvp19','')} | {c['dim1']} | {c['dim2']} | {c['dim3']} | "
            f"{c['dim4']} | {c['fail_tag']} | {note} |"
        )
    lines.append("")
    lines.append("## 逐封说明")
    lines.append("")
    for c in cases:
        lines.extend(
            [
                f"### {c['scenario_id']} · {c.get('mail_id','')}",
                "",
                f"- ① {c['dim1']} — {c['dim1_note']}",
                f"- ② {c['dim2']} — {c['dim2_note']}",
                f"- ③ {c['dim3']} — {c['dim3_note']}",
                f"- ④ **{c['dim4']}** · fail_tag={c['fail_tag']} · top1={c.get('top1')} · style={c.get('style')}",
                "",
            ]
        )
    return "\n".join(lines)


def apply_draft_to_xlsx(cases: list[dict], sheet: str, *, dry_run: bool) -> None:
    import openpyxl

    xlsx = EVAL / "22封真邮_四维评分表.xlsx"
    by_mail = {c.get("mail_id"): c for c in cases if c.get("mail_id")}
    wb = openpyxl.load_workbook(xlsx)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet}")
    ws = wb[sheet]
    applied = 0
    for r in range(5, 27):
        mail_id = ws.cell(r, 1).value
        if not mail_id:
            continue
        c = by_mail.get(str(mail_id))
        if not c:
            continue
        ws.cell(r, 5).value = c["dim1"]
        ws.cell(r, 6).value = c["dim2"]
        ws.cell(r, 7).value = c["dim3"]
        ws.cell(r, 8).value = c["dim4"] if c["dim4"] != "待人工" else None
        existing_i = ws.cell(r, 9).value
        tag = f"[草稿]{c['fail_tag']}"
        ws.cell(r, 9).value = tag if not existing_i else f"{existing_i}; {tag}"
        applied += 1
    if dry_run:
        print(f"DRY-RUN would apply {applied} rows to sheet {sheet!r}")
        return
    try:
        wb.save(xlsx)
        print(f"Applied draft E–I to {applied} rows in {sheet!r} (④待人工留空 · I 列带 [草稿] 前缀)")
    except PermissionError:
        out = xlsx.with_name(xlsx.stem + "_draft_scored.xlsx")
        wb.save(out)
        print(f"WARN xlsx locked — saved {out.name}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Draft 22-mail 4-dim scoring")
    ap.add_argument("--json", type=Path, default=EVAL / "cs_22mail_eval_round1.json")
    ap.add_argument(
        "--apply-sheet",
        default=None,
        help="Write draft to xlsx sheet (e.g. 复测评分(Wave1后)); ④待人工 leaves H empty",
    )
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    payload = load_json(args.json)
    scenarios = {s["id"]: s for s in load_json(MAP_PATH)["scenarios"]}
    cases = [draft_case(r, scenarios[r["scenario_id"]]) for r in payload["results"]]

    stem = args.json.stem.replace("cs_22mail_eval_", "")
    out_json = EVAL / f"scoring_draft_{stem}.json"
    out_md = EVAL / f"scoring_draft_{stem}.md"
    meta = {
        "source_json": args.json.name,
        "generated_at": payload.get("meta", {}).get("generated_at"),
        "round": payload.get("meta", {}).get("round"),
    }
    out_json.write_text(
        json.dumps({"meta": {**meta, "disclaimer": "NOT official scores"}, "cases": cases}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    out_md.write_text(
        render_md(cases, meta, title=f"22 封真邮 · 四维评分草稿 · {stem}"),
        encoding="utf-8",
    )
    mvp = [c for c in cases if c.get("mvp19") == "是"]
    pass_strict = sum(1 for c in mvp if c["dim4"] == "通过")
    print(f"Wrote {out_json.name} + {out_md.name}")
    print(f"MVP19 draft: {pass_strict}/{len(mvp)} strict pass (human must confirm)")

    if args.apply_sheet:
        apply_draft_to_xlsx(cases, args.apply_sheet, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
