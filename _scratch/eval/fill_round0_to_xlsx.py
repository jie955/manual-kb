#!/usr/bin/env python3
"""Fill scoring sheet from cs_22mail_eval_roundN.json (system output only, not ①②③④)."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

EVAL = Path(__file__).resolve().parent
ROOT = EVAL.parents[1]
DEFAULT_JSON = EVAL / "cs_22mail_eval_round0.json"
XLSX_PATH = EVAL / "22封真邮_四维评分表.xlsx"
MAP_PATH = EVAL / "cs_email_query_map.json"
DATA_START = 5
DATA_END = 26
URL_COL = 12  # L · Reference链接


def summarize_email(text: str, limit: int = 120) -> str:
    one = " ".join(text.split())
    if len(one) <= limit:
        return one
    return one[: limit - 1] + "…"


def format_images(images: list) -> str:
    if not images:
        return "(无)"
    seen: list[str] = []
    for img in images:
        name = img.get("file") or "?"
        if name not in seen:
            seen.append(name)
    return ", ".join(seen)


def format_preflight(row: dict) -> str:
    audit = row.get("reply_audit") or {}
    parts = [
        f"truncated={'Y' if row.get('reply_truncated') else 'N'}",
        f"steps={audit.get('step_count', '—')}",
        f"proceed_ref={'Y' if audit.get('proceed_ref') else 'N'}",
        f"check_step_ref={'Y' if audit.get('check_step_ref') else 'N'}",
        f"glued={'Y' if audit.get('glued_steps') else 'N'}",
        f"type={audit.get('mail_type', '—')}",
    ]
    inc = row.get("reply_incomplete_reason")
    if inc:
        parts.append(f"incomplete={inc}")
    return "[预检] " + " · ".join(parts)


def build_remark(row: dict) -> str:
    parts = [format_preflight(row)]
    parts.extend(
        [
            f"cs_id={row['scenario_id']}",
            f"lib={row['matched_library']}({row['routing_method']})",
            f"top1={row.get('top1_group') or '—'}",
            f"top3={','.join(row.get('top3_groups') or [])}",
        ]
    )
    if row.get("top1_hit") is False:
        parts.append("top1_hit=MISS")
    elif row.get("top1_hit") is True:
        parts.append("top1_hit=OK")
    parts.append(f"images={format_images(row.get('images_used') or [])}")
    if row.get("style"):
        parts.append(f"style={row['style'].get('family_id')}")
    reply = (row.get("generated_reply_en") or "").strip()
    if reply:
        parts.append("")
        parts.append("--- system reply ---")
        parts.append(reply)
    return "\n".join(parts)


def model_label(scenario: dict | None, matched_library: str) -> str:
    if scenario:
        pl = scenario.get("product_line")
        if pl:
            return str(pl)
    return matched_library


def corpus_path(scenario: dict | None) -> Path | None:
    if not scenario:
        return None
    rel = scenario.get("file")
    if not rel:
        return None
    path = (ROOT / rel).resolve()
    return path if path.is_file() else None


def set_reference_link(cell, path: Path) -> None:
    cell.value = "打开 Reference"
    cell.hyperlink = path.as_uri()
    cell.font = Font(color="0563C1", underline="single")
    cell.alignment = Alignment(vertical="top")


def main() -> None:
    ap = argparse.ArgumentParser(description="Fill xlsx from batch eval JSON")
    ap.add_argument("--json", type=Path, default=DEFAULT_JSON, help="eval JSON path")
    ap.add_argument(
        "--sheet",
        default=None,
        help="worksheet name (default: 基线评分 for round0, 复测评分 for round1)",
    )
    args = ap.parse_args()

    json_path = args.json
    round_name = json_path.stem
    sheet = args.sheet
    if sheet is None:
        if "round1" in round_name:
            sheet = "复测评分(Wave1后)"
        else:
            sheet = "基线评分(修复前)"

    payload = json.loads(json_path.read_text(encoding="utf-8"))
    scenarios = {s["id"]: s for s in json.loads(MAP_PATH.read_text(encoding="utf-8"))["scenarios"]}
    by_mail = {r["mail_id"]: r for r in payload["results"]}

    wb = openpyxl.load_workbook(XLSX_PATH)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet!r} in {XLSX_PATH.name}")
    ws = wb[sheet]
    if ws.cell(4, URL_COL).value != "Reference链接":
        ws.cell(4, URL_COL).value = "Reference链接"
        ws.column_dimensions["L"].width = 16

    filled = 0
    for r in range(DATA_START, DATA_END + 1):
        mail_id = ws.cell(r, 1).value
        if not mail_id:
            continue
        row = by_mail.get(str(mail_id))
        if not row:
            print(f"WARN no json for {mail_id}")
            continue
        scen = scenarios.get(row["scenario_id"])
        ws.cell(r, 2).value = row["mvp19"]  # B 计入MVP19
        ws.cell(r, 3).value = model_label(scen, row["matched_library"])  # C 机型
        ws.cell(r, 4).value = summarize_email(row.get("raw_email") or "")  # D 问题摘要
        path = corpus_path(scen)
        if path:
            set_reference_link(ws.cell(r, URL_COL), path)  # L Reference链接
        # E-H ①②③④ left for human
        ws.cell(r, 10).value = build_remark(row)  # J 备注
        ws.cell(r, 10).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        filled += 1

    try:
        wb.save(XLSX_PATH)
        out = XLSX_PATH
    except PermissionError:
        out = XLSX_PATH.with_name(XLSX_PATH.stem + "_filled.xlsx")
        wb.save(out)
        print(f"WARN {XLSX_PATH.name} is open — saved to {out.name}; close Excel and re-run to overwrite.")
        return
    print(f"Filled {filled} rows in {sheet} from {json_path.name}")


if __name__ == "__main__":
    main()
