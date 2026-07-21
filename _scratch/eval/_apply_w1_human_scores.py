#!/usr/bin/env python3
"""Apply human ①–④ scores to xlsx (Round 1 · 复测评分 · P0/P1/P2 定案)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import openpyxl

EVAL = Path(__file__).resolve().parent
XLSX = EVAL / "22封真邮_四维评分表.xlsx"
SHEET = "复测评分(Wave1后)"
SCORER = "agent-p0p1p2-2026-07-12"
SCORE_DATE = date.today().isoformat()

# Human scores · Round 1 · cs_22mail_eval_round1.json · P0/P1/P2 semantic review
HUMAN: dict[str, dict] = {
    "MAIL-01": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "P1①对·qa_011；DIP#3/4#5# instant short/限位短接；image_007配排查步",
    },
    "MAIL-02": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "top1=qa_037∉acceptable；缺11#/12#与限位短接",
    },
    "MAIL-03": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "P2·qa_010；DIP#5+4#5# short；image_007 batch重复·配板图可接受",
    },
    "MAIL-04": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "无需图",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "presales",
    },
    "MAIL-05": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "presales·image_007 caption空",
    },
    "MAIL-06": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "presales weak links",
    },
    "MAIL-07": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "presales",
    },
    "MAIL-08": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "TC148·不计MVP19",
    },
    "MAIL-09": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "top1=qa_031∉acceptable",
    },
    "MAIL-10": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "配对相关",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "TC148·不计MVP19",
    },
    "MAIL-11": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "无需图",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "presales weak links",
    },
    "MAIL-12": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "P0·qa_019∉{qa_020,qa_021}；真缺FORCE/SOFT_STOP/11#12#",
    },
    "MAIL-13": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "无需图",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "P2·qa_022；断电FORCE/SOFT STOP齐",
    },
    "MAIL-14": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "P0/P1·应qa_018磁环路径；qa_019限位 playbook偏",
    },
    "MAIL-15": {
        "dim1": "对",
        "dim2": "直接可发",
        "dim3": "无需图",
        "dim4": "通过",
        "fail_tag": "—",
        "note": "P2·Power off→FORCE max 四步对齐",
    },
    "MAIL-16": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "P0·ref_keys假阳性；真缺DIP#5/4#5#/单臂测；image_001 fuse OK",
    },
    "MAIL-17": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "P2·100%窄词表假阳性；缺fuse/4#5#/11#12#；image_019限位短接OK",
    },
    "MAIL-18": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "缺DIP#3/限位短接等",
    },
    "MAIL-19": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "无需图",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "P0·qa_015∉acceptable；真缺limit B/遥控学码",
    },
    "MAIL-20": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "无需图",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "P1①对·ad5s；top1=qa_040同窗≠qa_016；真缺limit B重确认",
    },
    "MAIL-21": {
        "dim1": "错",
        "dim2": "小改可发",
        "dim3": "配对相关",
        "dim4": "不通过",
        "fail_tag": "复合",
        "note": "P0·qa_033≠qa_001；真缺11#/12#+4#5# instant short",
    },
    "MAIL-22": {
        "dim1": "对",
        "dim2": "小改可发",
        "dim3": "无需图",
        "dim4": "不通过",
        "fail_tag": "生成",
        "note": "P0/P1·qa_040∈expected；真缺limit A/单臂隔离",
    },
}


def main() -> None:
    import subprocess
    import sys

    subprocess.run(
        [
            sys.executable,
            str(EVAL / "fill_round0_to_xlsx.py"),
            "--json",
            str(EVAL / "cs_22mail_eval_round1.json"),
            "--sheet",
            SHEET,
        ],
        check=True,
        cwd=EVAL.parents[1],
    )

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    applied = 0
    mvp_pass = 0
    mvp_total = 0
    for r in range(5, 27):
        mail_id = str(ws.cell(r, 1).value or "")
        if mail_id not in HUMAN:
            continue
        h = HUMAN[mail_id]
        ws.cell(r, 5).value = h["dim1"]
        ws.cell(r, 6).value = h["dim2"]
        ws.cell(r, 7).value = h["dim3"]
        ws.cell(r, 8).value = h["dim4"]
        ws.cell(r, 9).value = h["fail_tag"]
        ws.cell(r, 11).value = f"{SCORER} · {SCORE_DATE} · {h['note']}"
        applied += 1
        in_mvp = str(ws.cell(r, 2).value or "").strip() in ("是", "Y", "yes")
        if in_mvp:
            mvp_total += 1
            if h["dim4"] == "通过":
                mvp_pass += 1

    summary = {
        "sheet": SHEET,
        "source_json": "cs_22mail_eval_round1.json",
        "scorer": SCORER,
        "date": SCORE_DATE,
        "basis": "P0/P1/P2 semantic review 2026-07-10..12",
        "applied_rows": applied,
        "mvp19_pass": mvp_pass,
        "mvp19_total": mvp_total,
        "full22_pass": sum(1 for h in HUMAN.values() if h["dim4"] == "通过"),
        "scores": HUMAN,
    }
    out_json = EVAL / "scoring_human_round1.json"
    out_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    md = render_md(summary)
    (EVAL / "scoring_human_round1.md").write_text(md, encoding="utf-8")
    try:
        wb.save(XLSX)
        saved = str(XLSX)
    except PermissionError:
        saved = str(XLSX.with_name(XLSX.stem + "_human_scored.xlsx"))
        wb.save(saved)

    print(f"Saved xlsx: {saved}")
    print(f"MVP19 ④通过: {mvp_pass}/{mvp_total} (目标≥16/19)")
    print(f"全22封 ④通过: {summary['full22_pass']}/22")
    print(f"Wrote {out_json.name} + scoring_human_round1.md")


def render_md(summary: dict) -> str:
    lines = [
        "# 22 封真邮 · 人工四维评分 · Round 1（P0/P1/P2 定案）",
        "",
        f"**日期**：{summary['date']} · **评分人**：{summary['scorer']}",
        f"**依据**：{summary.get('basis', 'P0/P1/P2')}",
        f"**来源**：`{summary['source_json']}` · sheet `{summary['sheet']}`",
        "**路径**：07-08 无补丁 Round 1（≠ Round 1e）",
        "",
        "## 汇总",
        "",
        f"| 指标 | 值 |",
        f"| --- | ---: |",
        f"| **MVP19 ④通过（发链线）** | **{summary['mvp19_pass']}/{summary['mvp19_total']}** |",
        f"| 目标 | ≥16/19 |",
        f"| 全 22 封 ④通过 | {summary['full22_pass']}/22 |",
        "",
        "## 全表",
        "",
        "| mail | ① | ② | ③ | ④ | fail_tag | 备注 |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for mail_id in sorted(HUMAN.keys(), key=lambda x: int(x.replace("MAIL-", ""))):
        h = HUMAN[mail_id]
        lines.append(
            f"| {mail_id} | {h['dim1']} | {h['dim2']} | {h['dim3']} | {h['dim4']} | "
            f"{h['fail_tag']} | {h['note']} |"
        )
    lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    main()
